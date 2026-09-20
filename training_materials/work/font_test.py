from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

fonts = ['Arial Unicode MS', 'Heiti SC', 'Hiragino Sans GB', 'STHeiti', 'SimSong', 'FZHei-B01', 'Apple LiGothic']
d = Document()
for f in fonts:
    p = d.add_paragraph()
    r = p.add_run(f'{f}: 天马智擎平台使用培训 首页对话 工作台 知识中心')
    r.font.name = f
    r.font.size = Pt(18)
    rp = r._element.get_or_add_rPr()
    rp.rFonts.set(qn('w:ascii'), f)
    rp.rFonts.set(qn('w:hAnsi'), f)
    rp.rFonts.set(qn('w:eastAsia'), f)
d.save('/Users/richelleshi/workspace/portal/training_materials/work/font_test.docx')
