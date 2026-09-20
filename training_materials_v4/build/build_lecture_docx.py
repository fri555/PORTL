from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/Users/richelleshi/workspace/portal/training_materials_v4')
DATA = json.loads((ROOT / 'build/lecture_data.json').read_text(encoding='utf-8'))
RENDER = ROOT / 'build/pptx/render'
OUT = ROOT / '天马智擎平台培训讲稿_逐页对应版.docx'

ORANGE = 'FF6600'
INK = '272727'
MUTED = '707070'
PALE = 'FFF3E8'
LINE = 'E6E6E6'
FONT = 'Arial Unicode MS'


def set_fonts(run, size=11, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), FONT)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def borders(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_el = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:color'), color)
        borders_el.append(el)
    tc_pr.append(borders_el)


def set_cell_margin(cell, top=180, start=220, bottom=180, end=220):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = Inches(0.38)
section.bottom_margin = Inches(0.38)
section.left_margin = Inches(0.45)
section.right_margin = Inches(0.45)
section.header_distance = Inches(0.18)
section.footer_distance = Inches(0.18)

normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.28

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_fonts(header.add_run('TIANMA  |  天马智擎内部培训 · 逐页讲稿'), 8.5, True, ORANGE)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_fonts(footer.add_run('内部培训材料  ·  重要输出请人工复核'), 8, False, MUTED)

for idx, item in enumerate(DATA):
    if idx:
        doc.add_page_break()

    top = doc.add_table(rows=1, cols=2)
    top.autofit = False
    top.columns[0].width = Inches(8.9)
    top.columns[1].width = Inches(1.7)
    c0, c1 = top.rows[0].cells
    c0.width = Inches(8.9)
    c1.width = Inches(1.7)
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c0.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_fonts(p.add_run(item['title']), 20, True, INK)
    p2 = c1.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_fonts(p2.add_run(f"第 {item['page']:02d} / {len(DATA):02d} 页"), 10, True, ORANGE)

    if item.get('subtitle'):
        sub = doc.add_paragraph()
        sub.paragraph_format.space_before = Pt(0)
        sub.paragraph_format.space_after = Pt(7)
        set_fonts(sub.add_run(item['subtitle']), 9.5, True, ORANGE)

    body = doc.add_table(rows=1, cols=2)
    body.autofit = False
    left, right = body.rows[0].cells
    left.width = Inches(6.55)
    right.width = Inches(4.05)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margin(left, 100, 100, 100, 100)
    set_cell_margin(right, 260, 300, 220, 300)
    borders(left)
    borders(right, ORANGE)
    shade(right, PALE)

    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_after = Pt(0)
    run = lp.add_run()
    run.add_picture(str(RENDER / f"slide-{item['page']:02d}.png"), width=Inches(6.25))

    rp = right.paragraphs[0]
    rp.paragraph_format.space_after = Pt(10)
    set_fonts(rp.add_run('讲师讲稿'), 12, True, ORANGE)
    speech = right.add_paragraph()
    speech.paragraph_format.space_after = Pt(12)
    speech.paragraph_format.line_spacing = 1.38
    set_fonts(speech.add_run(item['note']), 11.5, False, INK)

    cue = right.add_paragraph()
    cue.paragraph_format.space_before = Pt(4)
    cue.paragraph_format.space_after = Pt(2)
    set_fonts(cue.add_run('现场提示'), 10, True, ORANGE)
    cue_text = '按截图中的红框和箭头依次演示，完成一步后停顿确认。' if item.get('source') else '本页作为转场或总结，不展开重复说明。'
    cue2 = right.add_paragraph()
    cue2.paragraph_format.space_after = Pt(0)
    set_fonts(cue2.add_run(cue_text), 9.5, False, MUTED)

doc.save(OUT)
print(OUT)
