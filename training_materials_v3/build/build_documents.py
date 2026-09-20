from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("training_materials_v3").resolve()
ORANGE = "FF6A00"
INK = "222222"
MUTED = "666666"
LIGHT = "FFF2E8"
FONT = "Heiti SC"

def set_rfonts(rpr):
    rfonts = rpr.rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_run(run, size=10.5, bold=False, color=INK):
    run.font.name = FONT
    set_rfonts(run._element.get_or_add_rPr())
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run(run, 9, False, MUTED)

def new_decimal_num_id(doc):
    numbering = doc.part.numbering_part.element
    ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(ids or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "0")
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id

def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    paragraph.paragraph_format.left_indent = Inches(0.38)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)

def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    set_rfonts(normal._element.rPr)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in [
        ("Title", 28, 0, 10, ORANGE),
        ("Subtitle", 13, 0, 10, MUTED),
        ("Heading 1", 19, 18, 10, ORANGE),
        ("Heading 2", 15, 14, 7, INK),
        ("Heading 3", 12, 10, 5, ORANGE),
        ("Heading 4", 11, 8, 4, INK),
    ]:
        st = styles[name]
        st.font.name = FONT
        set_rfonts(st._element.rPr)
        st.font.size = Pt(size)
        st.font.bold = name != "Subtitle"
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        st = styles[name]
        st.font.name = FONT
        set_rfonts(st._element.rPr)
        st.font.size = Pt(10.5)
        st.paragraph_format.left_indent = Inches(0.38)
        st.paragraph_format.first_line_indent = Inches(-0.18)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "TIANMA  |  天马智擎内部培训"
    for run in header.runs:
        set_run(run, 8.5, True, ORANGE)
    footer = section.footer.paragraphs[0]
    footer.add_run("内部培训材料  ·  请按数据规范使用    ")
    set_run(footer.runs[0], 8.5, False, MUTED)
    add_page_number(footer)

def add_rich_text(paragraph, text, size=10.5, color=INK):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run(paragraph.add_run(part[2:-2]), size, True, ORANGE)
        elif part.startswith("`") and part.endswith("`"):
            set_run(paragraph.add_run(part[1:-1]), size, True, MUTED)
        else:
            set_run(paragraph.add_run(part), size, False, color)

def add_cover(doc, title, subtitle):
    doc.add_paragraph().paragraph_format.space_after = Pt(80)
    kicker = doc.add_paragraph()
    add_rich_text(kicker, "TIANMA · INTERNAL TRAINING", 10, ORANGE)
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    add_rich_text(p, title, 28, ORANGE)
    s = doc.add_paragraph(style="Subtitle")
    add_rich_text(s, subtitle, 13, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(35)
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(2.15), Inches(2.15), Inches(2.15)]
    for idx, (label, value) in enumerate([
        ("适用版本", "V1 内测版"),
        ("环境依据", "2026-08-08 正式环境"),
        ("视觉规范", "天马橙 · 1920×1080"),
    ]):
        cell = table.cell(0, idx)
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT)
        p1 = cell.paragraphs[0]
        add_rich_text(p1, label, 9, MUTED)
        p2 = cell.add_paragraph()
        add_rich_text(p2, value, 10.5, ORANGE)
    doc.add_paragraph().paragraph_format.space_after = Pt(125)
    note = doc.add_paragraph()
    add_rich_text(note, "选对模式 · 给足材料 · 说清任务 · 人工复核 · 有效反馈", 12, INK)
    doc.add_page_break()

def add_toc(doc, lines):
    doc.add_heading("目录", level=1)
    for level, title in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0 if level == 1 else 0.28)
        add_rich_text(p, title, 11 if level == 1 else 10, ORANGE if level == 1 else INK)
    doc.add_page_break()

SCRIPT_IMAGE_MAP = {
    "2.1 首页导览：布局与关键入口": "01_home_overview.png",
    "2.1.4 搜索和打开历史会话": "10_conversation_search.png",
    "2.1.7 日常办公和专家模式": "02_mode_positioning.png",
    "2.1.9 上传本地文件": "05_local_attachment_limits.png",
    "2.1.10 从知识中心引用文件": "06_knowledge_picker.png",
    "2.2 知识中心：知识库建设与关联配置": "11_knowledge_center.png",
    "2.3 工作台：内部系统统一入口": "13_workbench.png",
    "3.3 数据分析：从问题定义到结果沉淀": "04_data_analysis_agent.png",
    "3.3.5 第四步：选择数据分析专家并挂载文件": "07_data_file_selected.png",
    "3.3.6 第五步：用六要素写分析任务": "08_data_analysis_prompt.png",
    "3.3.7 第六步：查看执行过程和异常状态": "09_data_analysis_processing.png",
}

def parse_markdown(path, output, title, subtitle, insert_script_images=False):
    lines = path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    add_cover(doc, title, subtitle)
    toc = []
    for line in lines:
        m = re.match(r"^(#{1,2})\s+(.+)$", line)
        if m and m.group(2) != title:
            toc.append((len(m.group(1)), m.group(2)))
    add_toc(doc, toc[:32])

    first_h1 = True
    image_re = re.compile(r"!\[(.*?)\]\((.*?)\)")
    current_num_id = None
    for raw in lines:
        line = raw.rstrip()
        if not line or line == "---":
            current_num_id = None
            continue
        if line == f"# {title}":
            continue
        image_match = image_re.fullmatch(line)
        if image_match:
            current_num_id = None
            image_path = (path.parent / image_match.group(2)).resolve()
            doc.add_picture(str(image_path), width=Inches(6.6))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(cap, f"图：{image_match.group(1)}", 8.5, MUTED)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            current_num_id = None
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1 and not first_h1:
                doc.add_page_break()
            first_h1 = False if level == 1 else first_h1
            p = doc.add_heading(level=min(level, 4))
            add_rich_text(p, text, [0, 19, 15, 12, 11][level], ORANGE if level in (1, 3) else INK)
            if insert_script_images and text in SCRIPT_IMAGE_MAP:
                img = ROOT / "screenshots/annotated" / SCRIPT_IMAGE_MAP[text]
                doc.add_picture(str(img), width=Inches(6.6))
            continue
        if line.startswith("> "):
            current_num_id = None
            table = doc.add_table(rows=1, cols=1)
            table.autofit = False
            cell = table.cell(0, 0)
            cell.width = Inches(6.6)
            set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            add_rich_text(p, line[2:], 10, INK)
            continue
        if re.match(r"^\d+\.\s+", line):
            if current_num_id is None:
                current_num_id = new_decimal_num_id(doc)
            p = doc.add_paragraph()
            apply_numbering(p, current_num_id)
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", line))
            continue
        if line.startswith("- "):
            current_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
            continue
        p = doc.add_paragraph()
        current_num_id = None
        add_rich_text(p, line)

    doc.core_properties.title = title
    doc.core_properties.subject = "天马智擎内部培训"
    doc.save(output)

parse_markdown(
    ROOT / "天马智擎平台系统使用说明_正式版.md",
    ROOT / "天马智擎平台系统使用说明_正式版.docx",
    "天马智擎平台系统使用说明",
    "平台定位 · 首页 · 知识中心 · 工作台 · 问题反馈 · 常见问题",
)
parse_markdown(
    ROOT / "天马智擎平台培训讲稿_正式版.md",
    ROOT / "天马智擎平台培训讲稿_正式版.docx",
    "天马智擎平台培训讲稿",
    "四章结构逐字讲解版 · 含现场操作与能力边界",
    insert_script_images=True,
)
